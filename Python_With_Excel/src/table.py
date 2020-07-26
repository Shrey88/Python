""" *
    *   we are going to create a simple table which is iterable by the reader as well as importing images
    *
    *   Image manipulation is not done solely through openpyxl and you may need to install a library called Pillow
    * """
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image

wb = load_workbook(".\\Exercise Files\\Pie.xlsx")
ws = wb.active

#   we will create a table object with display name of Table1
tab = Table(displayName='Table1', ref='A1:B5')

#   now the table knows what its called and where to find the data and it needs to know how to format itself.
#   from here we will use table style information
style = TableStyleInfo(name='TableStyleMedium9', showFirstColumn=False, showLastColumn=False, showRowStripes=True,
                       showColumnStripes=True)

#   set the tablestyleinfo property to style we just created.
tab.tableStyleInfo = style

#   add the table to worksheet
ws.add_table(tab)

img = Image(".\\Exercise Files\\mountain.jpg")
img.height = img.height * .25
img.width = img.width * .25

ws.add_image(img, 'C1')

#   image property allows us to access us height and width
wb.save(".\\Exercise Files\\table.xlsx")