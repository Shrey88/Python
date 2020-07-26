""" *
    *   when working with reports, a graphical interface is important to keep in mind. Openpyxl does not support
    *   features to work with the pivot tables, but we can format our reports to look nice after we've analyzed our data
    *   in pandas
    * """
from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart, Series, Reference

wb = load_workbook(".\\Exercise Files\\crime_report.xlsx")
ws = wb.active

chart = BarChart()

data = Reference(ws, min_row=10, min_col=1, max_col=13, max_row=13)
labels = Reference(ws, min_row=8, min_col=2, max_row=8, max_col=13)
chart.add_data(data, from_rows=True, titles_from_data=True)
chart.set_categories(labels)
chart.title= "Counterfeit Crimes by District"
chart.height = chart.height * .75
chart.width = chart.width * .55
ws.add_chart(chart, "B14")
wb.save(".\\Exercise Files\\lines.xlsx")